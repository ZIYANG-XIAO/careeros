"""Evidence-grounded materials: tailoring, drafts, interview prep."""

from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from conftest import make_integration_client, make_test_settings
from docx import Document
from fastapi.testclient import TestClient
from test_job_intelligence import (
    FakeJobParser,
    FakeResumeParser,
    create_job,
    create_resume_and_evidence,
)

from careeros_api.job_intelligence_schemas import EvidenceInput, RequirementInput
from careeros_api.main import create_app
from careeros_api.materials_schemas import (
    DraftResult,
    InterviewPrepBatch,
    PrepGap,
    PrepQuestion,
    PrepStory,
    TailoringBatch,
    TailoringSuggestion,
    apply_tailoring,
    introduces_unsupported_numbers,
)

RESUME_LINE = "Backend engineer with Python experience"


def _by(items: list[EvidenceInput], text: str) -> UUID:
    return next(item.revision_id for item in items if text in item.content)


def _req(items: list[RequirementInput], value: str) -> UUID:
    return next(item.id for item in items if item.normalized_value == value)


class FakeMaterialsWriter:
    model = "fake-materials-writer"

    def __init__(self, draft_text: str = "I built Kafka pipelines at work.") -> None:
        self.draft_text = draft_text

    async def suggest_tailoring(
        self,
        requirements: list[RequirementInput],
        evidence: list[EvidenceInput],
        resume_text: str,
    ) -> TailoringBatch:
        kafka_req = _req(requirements, "Kafka")
        kafka_ev = _by(evidence, "Kafka")
        base = {"requirement_ids": [kafka_req], "rationale": "Surfaces confirmed Kafka work."}
        return TailoringBatch(
            suggestions=[
                TailoringSuggestion(
                    kind="REWRITE",
                    original_text=RESUME_LINE,
                    insert_after=None,
                    suggested_text="Backend engineer with Python and Kafka experience",
                    evidence_revision_ids=[kafka_ev],
                    **base,
                ),
                # Anchor not in the resume.
                TailoringSuggestion(
                    kind="REWRITE",
                    original_text="Led a team of engineers",
                    insert_after=None,
                    suggested_text="Led Kafka engineers",
                    evidence_revision_ids=[kafka_ev],
                    **base,
                ),
                # Invents a number.
                TailoringSuggestion(
                    kind="INSERT",
                    original_text=None,
                    insert_after=RESUME_LINE,
                    suggested_text="Scaled Kafka to 40 brokers",
                    evidence_revision_ids=[kafka_ev],
                    **base,
                ),
                # Cites evidence outside the candidate set.
                TailoringSuggestion(
                    kind="INSERT",
                    original_text=None,
                    insert_after=RESUME_LINE,
                    suggested_text="Operated Kafka",
                    evidence_revision_ids=[uuid4()],
                    **base,
                ),
            ]
        )

    async def draft_answer(
        self,
        kind: str,
        question: str | None,
        job_text: str,
        evidence: list[EvidenceInput],
    ) -> DraftResult:
        return DraftResult(
            text=self.draft_text,
            evidence_revision_ids=[_by(evidence, "Kafka")],
            unsupported_points=["GCP production experience"],
        )

    async def prepare_interview(
        self,
        requirements: list[RequirementInput],
        evidence: list[EvidenceInput],
        job_text: str,
    ) -> InterviewPrepBatch:
        kafka_req = _req(requirements, "Kafka")
        return InterviewPrepBatch(
            stories=[
                PrepStory(
                    requirement_ids=[kafka_req],
                    evidence_revision_ids=[_by(evidence, "Kafka")],
                    prompt="Tell the story of your Kafka work.",
                ),
                # No evidence: a story must rest on something real.
                PrepStory(requirement_ids=[kafka_req], evidence_revision_ids=[], prompt="x"),
            ],
            questions=[
                PrepQuestion(
                    question="How do you handle consumer lag?",
                    requirement_ids=[kafka_req],
                    evidence_revision_ids=[],
                )
            ],
            gaps=[
                PrepGap(
                    requirement_id=_req(requirements, "GCP"),
                    honest_framing="Say you have not run GCP in production; mention adjacent work.",
                ),
                PrepGap(requirement_id=uuid4(), honest_framing="unknown requirement"),
            ],
        )


def test_number_guard_and_edit_application() -> None:
    assert introduces_unsupported_numbers("Cut latency 40%", "latency work")
    assert not introduces_unsupported_numbers("Cut latency 40%", "reduced latency by 40%")
    assert not introduces_unsupported_numbers("Served 1,200 users", "1200 users")
    text = apply_tailoring(
        "A\nB",
        [("REWRITE", "A", None, "A2"), ("INSERT", None, "B", "C")],
    )
    assert text == "A2\nB\nC"
    with pytest.raises(ValueError):
        apply_tailoring("A", [("REWRITE", "missing", None, "x")])


@pytest.mark.asyncio
async def test_tailoring_keeps_only_grounded_suggestions_and_creates_a_version(
    tmp_path: Path,
) -> None:
    async with make_integration_client(
        resume_parser=FakeResumeParser(),
        job_parser=FakeJobParser(),
        materials_writer=FakeMaterialsWriter(),
        resume_storage_dir=str(tmp_path),
    ) as client:
        version_id, _ = await create_resume_and_evidence(client)
        job_id = await create_job(client)

        response = await client.post(
            f"/jobs/{job_id}/tailoring", json={"resume_version_id": version_id}
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["dropped_count"] == 3
        [suggestion] = body["suggestions"]
        assert suggestion["original_text"] == RESUME_LINE
        assert suggestion["evidence"][0]["content"] == "Kafka"
        assert suggestion["requirements"][0]["description"] == "Kafka experience required."

        created = await client.post(
            "/resumes/tailored",
            json={
                "base_resume_version_id": version_id,
                "job_id": job_id,
                "edits": [
                    {
                        "kind": "REWRITE",
                        "original_text": RESUME_LINE,
                        "text": suggestion["suggested_text"],
                    }
                ],
            },
        )
        assert created.status_code == 201, created.text
        new_version = created.json()
        assert new_version["version_number"] == 2
        assert new_version["original_filename"].endswith(".docx")

        resume = (await client.get(f"/resumes/{new_version['resume_id']}")).json()
        versions = {item["id"]: item for item in resume["versions"]}
        assert versions[version_id]["extracted_text"].startswith(RESUME_LINE)
        tailored = versions[new_version["resume_version_id"]]
        assert tailored["source"] == "TAILORED"
        assert tailored["source_version_id"] == version_id
        assert "Python and Kafka experience" in tailored["extracted_text"]

        file = await client.get(f"/resume-versions/{new_version['resume_version_id']}/file")
        assert file.status_code == 200
        paragraphs = [p.text for p in Document(BytesIO(file.content)).paragraphs]
        assert "Backend engineer with Python and Kafka experience" in paragraphs

        stale = await client.post(
            "/resumes/tailored",
            json={
                "base_resume_version_id": version_id,
                "job_id": job_id,
                "edits": [{"kind": "REWRITE", "original_text": "not there", "text": "x"}],
            },
        )
        assert stale.status_code == 422
        assert stale.json()["detail"]["code"] == "EDIT_ANCHOR_MISSING"


@pytest.mark.asyncio
async def test_drafts_are_grounded_and_refuse_sensitive_questions(tmp_path: Path) -> None:
    async with make_integration_client(
        resume_parser=FakeResumeParser(),
        job_parser=FakeJobParser(),
        materials_writer=FakeMaterialsWriter(),
        resume_storage_dir=str(tmp_path),
    ) as client:
        no_evidence = await client.post(
            "/drafts", json={"kind": "OPEN_QUESTION", "question": "Why us?"}
        )
        assert no_evidence.status_code == 409

        await create_resume_and_evidence(client)
        job_id = await create_job(client)
        draft = await client.post("/drafts", json={"kind": "COVER_LETTER", "job_id": job_id})
        assert draft.status_code == 200, draft.text
        assert draft.json()["evidence"][0]["content"] == "Kafka"
        assert draft.json()["unsupported_points"] == ["GCP production experience"]

        sensitive = await client.post(
            "/drafts",
            json={"kind": "OPEN_QUESTION", "question": "Will you require visa sponsorship?"},
        )
        assert sensitive.status_code == 422
        assert sensitive.json()["detail"]["code"] == "SENSITIVE_QUESTION"

        no_job = await client.post("/drafts", json={"kind": "COVER_LETTER"})
        assert no_job.json()["detail"]["code"] == "JOB_REQUIRED"


@pytest.mark.asyncio
async def test_draft_with_an_invented_number_is_rejected(tmp_path: Path) -> None:
    async with make_integration_client(
        resume_parser=FakeResumeParser(),
        job_parser=FakeJobParser(),
        materials_writer=FakeMaterialsWriter(draft_text="I ran Kafka for 12 years."),
        resume_storage_dir=str(tmp_path),
    ) as client:
        await create_resume_and_evidence(client)
        response = await client.post(
            "/drafts", json={"kind": "OPEN_QUESTION", "question": "Describe a project."}
        )
        assert response.status_code == 422
        assert response.json()["detail"]["code"] == "DRAFT_NOT_GROUNDED"


@pytest.mark.asyncio
async def test_interview_prep_drops_ungrounded_items(tmp_path: Path) -> None:
    async with make_integration_client(
        resume_parser=FakeResumeParser(),
        job_parser=FakeJobParser(),
        materials_writer=FakeMaterialsWriter(),
        resume_storage_dir=str(tmp_path),
    ) as client:
        await create_resume_and_evidence(client)
        job_id = await create_job(client)
        response = await client.post(f"/jobs/{job_id}/interview-prep")
        assert response.status_code == 200, response.text
        body = response.json()
        assert [story["prompt"] for story in body["stories"]] == [
            "Tell the story of your Kafka work."
        ]
        assert body["questions"][0]["question"] == "How do you handle consumer lag?"
        assert [gap["requirement"]["description"] for gap in body["gaps"]] == [
            "GCP production experience required."
        ]
        assert body["dropped_count"] == 2


@pytest.mark.asyncio
async def test_tailoring_rejects_another_users_resume(tmp_path: Path) -> None:
    async with make_integration_client(
        resume_parser=FakeResumeParser(),
        job_parser=FakeJobParser(),
        materials_writer=FakeMaterialsWriter(),
        resume_storage_dir=str(tmp_path / "owner"),
    ) as owner:
        version_id, _ = await create_resume_and_evidence(owner)
        async with make_integration_client(
            job_parser=FakeJobParser(),
            materials_writer=FakeMaterialsWriter(),
            resume_storage_dir=str(tmp_path / "other"),
        ) as other:
            job_id = await create_job(other)
            response = await other.post(
                f"/jobs/{job_id}/tailoring", json={"resume_version_id": version_id}
            )
            assert response.status_code == 404
            created = await other.post(
                "/resumes/tailored",
                json={
                    "base_resume_version_id": version_id,
                    "job_id": job_id,
                    "edits": [{"kind": "REWRITE", "original_text": RESUME_LINE, "text": "x"}],
                },
            )
            assert created.status_code == 404


@pytest.mark.parametrize(
    "path",
    [f"/jobs/{uuid4()}/tailoring", "/drafts", f"/jobs/{uuid4()}/interview-prep"],
)
def test_materials_require_authentication(path: str) -> None:
    with TestClient(create_app(make_test_settings())) as client:
        assert client.post(path, json={}).status_code == 401


@pytest.mark.asyncio
async def test_cover_letter_file_renders_the_approved_text_as_docx() -> None:
    async with make_integration_client() as client:
        response = await client.post(
            "/drafts/cover-letter-file",
            json={
                "text": "Dear team,\nI build payment systems.",
                "filename_hint": "Stripe / letter",
            },
        )
        assert response.status_code == 200
        assert 'filename="Stripe  letter.docx"' in response.headers["content-disposition"]
        paragraphs = [p.text for p in Document(BytesIO(response.content)).paragraphs]
        assert paragraphs == ["Dear team,", "I build payment systems."]
