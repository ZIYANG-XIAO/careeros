"""Deterministic validation and text extraction for untrusted resume files."""

from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024


class ResumeFileError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def detect_resume_type(data: bytes) -> tuple[str, str]:
    if data.startswith(b"%PDF-"):
        return PDF_CONTENT_TYPE, ".pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                total_size = sum(entry.file_size for entry in archive.infolist())
        except BadZipFile as error:
            raise ResumeFileError("MALFORMED_FILE", "The uploaded file is malformed") from error
        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            raise ResumeFileError(
                "UNSUPPORTED_FILE_TYPE", "Only PDF and DOCX resumes are supported"
            )
        if total_size > MAX_DOCX_UNCOMPRESSED_BYTES or len(names) > 1000:
            raise ResumeFileError("MALFORMED_FILE", "The DOCX archive is not safe to process")
        return DOCX_CONTENT_TYPE, ".docx"
    raise ResumeFileError("UNSUPPORTED_FILE_TYPE", "Only PDF and DOCX resumes are supported")


def extract_resume_text(data: bytes, content_type: str) -> str:
    try:
        if content_type == PDF_CONTENT_TYPE:
            reader = PdfReader(BytesIO(data), strict=True)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif content_type == DOCX_CONTENT_TYPE:
            document = Document(BytesIO(data))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(blocks)
        else:
            raise ResumeFileError("UNSUPPORTED_FILE_TYPE", "Unsupported resume file type")
    except ResumeFileError:
        raise
    except Exception as error:
        raise ResumeFileError("EXTRACTION_FAILED", "Resume text could not be extracted") from error

    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
    if not normalized:
        raise ResumeFileError("EMPTY_DOCUMENT", "The resume does not contain extractable text")
    return normalized
